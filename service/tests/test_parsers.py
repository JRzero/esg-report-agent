from io import BytesIO
from openpyxl import Workbook
from app.integrations.parsers import parse_xlsx

def test_xlsx_anchor_preserves_cell():
    wb=Workbook(); ws=wb.active; ws.title='员工统计'; ws['A1']='员工总数'; ws['B1']=1287
    b=BytesIO(); wb.save(b)
    anchors=parse_xlsx(b.getvalue())
    got=[x for x in anchors if x.cell_range=='B1'][0]
    assert got.sheet_name=='员工统计' and got.raw_text=='1287'
