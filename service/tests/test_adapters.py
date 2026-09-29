from io import BytesIO

import httpx
import pytest
from pptx import Presentation

from app.ai.llm import LLMGateway
from app.ai.openviking import OpenVikingAdapter
from app.ai.schemas import SectionPlan
from app.integrations.parsers import parse_pptx
from app.integrations.storage import storage


@pytest.mark.asyncio
async def test_llm_structured_output_contract():
    async def handler(request: httpx.Request):
        assert request.url.path == "/chat/completions"
        assert request.headers["authorization"] == "Bearer test-key"
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "message": {
                            "content": '{"goal":"Use confirmed facts","recommended_structure":["A"]}'
                        }
                    }
                ]
            },
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        gateway = LLMGateway(
            client=client,
            base_url="http://llm.test",
            api_key="test-key",
            model="test-model",
        )
        result = await gateway.generate_structured("system", "user", SectionPlan)

    assert result.goal == "Use confirmed facts"
    assert result.recommended_structure == ["A"]


@pytest.mark.asyncio
async def test_openviking_http_contract():
    calls = []

    async def handler(request: httpx.Request):
        calls.append((request.method, request.url.path, request.headers.get("x-api-key")))
        if request.url.path == "/health":
            return httpx.Response(200, json={"healthy": True})
        if request.url.path.endswith("/temp_upload"):
            return httpx.Response(200, json={"result": {"temp_file_id": "temp-1"}})
        if request.url.path == "/api/v1/resources":
            return httpx.Response(
                200,
                json={"result": {"root_uri": "viking://resource/1", "task_id": "task-1"}},
            )
        if request.url.path == "/api/v1/search/find":
            return httpx.Response(
                200,
                json={
                    "result": {
                        "results": [
                            {
                                "uri": "viking://resource/1",
                                "content": "employee evidence",
                                "score": 0.9,
                            }
                        ]
                    }
                },
            )
        raise AssertionError(f"Unexpected OpenViking path: {request.url.path}")

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        adapter = OpenVikingAdapter(
            client=client,
            base_url="http://openviking.test",
            api_key="secret",
        )
        assert await adapter.healthcheck()
        resource = await adapter.add_bytes(
            "employee.txt",
            b"employee evidence",
            "viking://resources/projects/p/evidence/employee.txt",
        )
        found = await adapter.find(
            "employee",
            "viking://resources/projects/p/evidence",
        )

    assert resource["task_id"] == "task-1"
    assert found[0].content == "employee evidence"
    assert all(call[2] == "secret" for call in calls)


def test_pptx_parser_preserves_slide_locator():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    textbox = slide.shapes.add_textbox(0, 0, 1000000, 1000000)
    textbox.text = "ESG治理架构"
    stream = BytesIO()
    presentation.save(stream)

    anchors = parse_pptx(stream.getvalue())
    assert anchors[0].slide_number == 1
    assert anchors[0].raw_text == "ESG治理架构"


@pytest.mark.asyncio
async def test_minio_storage_round_trip():
    backend = storage()
    assert await backend.healthcheck()
    key = "acceptance/storage/round-trip.txt"
    await backend.put(key, b"evidence", "text/plain")
    assert await backend.exists(key)
    assert await backend.get(key) == b"evidence"
    assert await backend.signed_download_url(key)
    await backend.delete(key)
    assert not await backend.exists(key)
