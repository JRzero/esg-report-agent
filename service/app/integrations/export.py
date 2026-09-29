from io import BytesIO
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.modules.models import Report, ReportSection, ReportBlock

async def render_docx(session:AsyncSession, report_id):
    from docx import Document
    report=await session.get(Report,report_id)
    doc=Document(); doc.add_heading(report.title,0)
    sections=list((await session.scalars(select(ReportSection).where(ReportSection.report_id==report_id).order_by(ReportSection.sort_order))).all())
    for s in sections:
        doc.add_heading(s.title,level=max(1,min(s.level,9)))
        blocks=list((await session.scalars(select(ReportBlock).where(ReportBlock.section_id==s.id,ReportBlock.deleted_at.is_(None)).order_by(ReportBlock.sort_order))).all())
        for b in blocks:
            if b.block_type=='PARAGRAPH': doc.add_paragraph(b.current_content)
            elif b.block_type=='HEADING': doc.add_heading(b.current_content,level=min(s.level+1,9))
            else: doc.add_paragraph(b.current_content)
    out=BytesIO(); doc.save(out); return out.getvalue()
