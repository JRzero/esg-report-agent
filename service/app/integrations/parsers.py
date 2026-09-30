from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO
from pathlib import Path


@dataclass
class AnchorData:
    anchor_type: str
    raw_text: str
    page_start: int | None = None
    page_end: int | None = None
    sheet_name: str | None = None
    cell_range: str | None = None
    heading_path: list[str] | None = None
    paragraph_start: int | None = None
    slide_number: int | None = None
    bbox: dict | None = None
    metadata: dict = field(default_factory=dict)

    @property
    def content_hash(self):
        locator = (
            f"{self.page_start}|{self.sheet_name}|{self.cell_range}|"
            f"{self.paragraph_start}|{self.slide_number}"
        )
        return sha256((locator + "|" + self.raw_text).encode()).hexdigest()


def parse_xlsx(data: bytes) -> list[AnchorData]:
    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(data), data_only=False, read_only=True)
    out = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            values = [cell.value for cell in row]
            row_context = [str(value) for value in values if value is not None][:20]
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if not text:
                    continue
                out.append(
                    AnchorData(
                        "EXCEL_CELL",
                        text,
                        sheet_name=worksheet.title,
                        cell_range=cell.coordinate,
                        metadata={"row_context": row_context},
                    )
                )
    return out


def parse_docx(data: bytes) -> list[AnchorData]:
    from docx import Document

    document = Document(BytesIO(data))
    headings: list[str] = []
    out = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            try:
                level = int(paragraph.style.name.split()[-1])
            except (TypeError, ValueError):
                level = 1
            headings = headings[: level - 1] + [text]
        out.append(
            AnchorData(
                "DOCX_PARAGRAPH",
                text,
                heading_path=list(headings),
                paragraph_start=index,
            )
        )
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            values = [cell.text.strip() for cell in row.cells]
            for col_index, value in enumerate(values):
                if not value:
                    continue
                out.append(
                    AnchorData(
                        "DOCX_TABLE_CELL",
                        value,
                        metadata={
                            "table_index": table_index,
                            "row_index": row_index,
                            "column_index": col_index,
                            "row_context": values,
                        },
                    )
                )
    return out


def parse_pdf(data: bytes) -> list[AnchorData]:
    import fitz

    document = fitz.open(stream=data, filetype="pdf")
    out = []
    for page_number, page in enumerate(document):
        for block in page.get_text("blocks"):
            text = (block[4] or "").strip()
            if not text:
                continue
            out.append(
                AnchorData(
                    "PDF_TEXT",
                    text,
                    page_start=page_number + 1,
                    page_end=page_number + 1,
                    bbox={"x0": block[0], "y0": block[1], "x1": block[2], "y1": block[3]},
                )
            )
    return out


def parse_pptx(data: bytes) -> list[AnchorData]:
    from pptx import Presentation

    presentation = Presentation(BytesIO(data))
    out = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes):
            text = getattr(shape, "text", "")
            text = text.strip() if text else ""
            if not text:
                continue
            out.append(
                AnchorData(
                    "PPTX_TEXT",
                    text,
                    slide_number=slide_number,
                    metadata={"shape_index": shape_index},
                )
            )
    return out


def parse_plain_text(data: bytes) -> list[AnchorData]:
    text = data.decode("utf-8-sig", errors="replace")
    out = []
    for index, line in enumerate(text.splitlines()):
        value = line.strip()
        if value:
            out.append(AnchorData("PLAIN_TEXT", value, paragraph_start=index))
    return out or [AnchorData("PLAIN_TEXT", text)]


def parse_document(filename: str, data: bytes) -> list[AnchorData]:
    extension = Path(filename).suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        return parse_xlsx(data)
    if extension == ".docx":
        return parse_docx(data)
    if extension == ".pdf":
        return parse_pdf(data)
    if extension == ".pptx":
        return parse_pptx(data)
    if extension in {".txt", ".md", ".csv"}:
        return parse_plain_text(data)
    raise ValueError(f"Unsupported file type: {extension}")
